"""Inspeção rápida dos arquivos VAAT para entender estrutura antes de escrever o notebook."""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from docx import Document

BASE = Path(r"c:\Users\carlos.marchionatti\GitHub - STN\STN-gered\VAAT")

def sep(titulo: str) -> None:
    print("\n" + "=" * 80)
    print(titulo)
    print("=" * 80)

# ---------- 1. VAAT_18-08-2026 1.xlsx ----------
vaat_path = BASE / "VAAT_18-08-2026 1.xlsx"
sep(f"ARQUIVO: {vaat_path.name}")
wb = load_workbook(vaat_path, read_only=True, data_only=True)
print("Abas:", wb.sheetnames)
for aba in wb.sheetnames:
    sep(f"  Aba: {aba}")
    df = pd.read_excel(vaat_path, sheet_name=aba, nrows=15)
    print("Shape (primeiras 15):", df.shape)
    print("Colunas:", list(df.columns))
    print(df.head(15).to_string())
    df_full = pd.read_excel(vaat_path, sheet_name=aba)
    print(f"\nTotal linhas na aba: {len(df_full)}")

# ---------- 2. Planilha Consolidação 18agosto2026.xlsx ----------
cons_path = BASE / "Planilha Consolidação 18agosto2026.xlsx"
sep(f"ARQUIVO: {cons_path.name}")
wb = load_workbook(cons_path, read_only=True, data_only=True)
print("Abas:", wb.sheetnames)
for aba in wb.sheetnames:
    sep(f"  Aba: {aba}")
    try:
        df = pd.read_excel(cons_path, sheet_name=aba, nrows=20)
        print("Shape (primeiras 20):", df.shape)
        print("Colunas:", list(df.columns))
        print(df.head(20).to_string())
        df_full = pd.read_excel(cons_path, sheet_name=aba)
        print(f"\nTotal linhas na aba: {len(df_full)}")
    except Exception as e:
        print(f"Erro lendo aba: {e}")

# ---------- 3. Comunicado pendências 18agosto26.docx ----------
doc_path = BASE / "Comunicado pendências 18agosto26.docx"
sep(f"ARQUIVO: {doc_path.name}")
doc = Document(doc_path)
print(f"Total de parágrafos: {len(doc.paragraphs)}")
print(f"Total de tabelas: {len(doc.tables)}")
sep("Texto (parágrafos):")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i:03d}] {p.text}")
if doc.tables:
    sep("Tabelas:")
    for ti, t in enumerate(doc.tables):
        print(f"\n-- Tabela {ti} ({len(t.rows)} linhas x {len(t.columns)} colunas) --")
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                print(f"  [L{ri}C{ci}] {cell.text.strip()[:120]}")
